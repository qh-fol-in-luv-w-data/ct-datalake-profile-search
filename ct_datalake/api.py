import frappe
import os
import json
import uuid
import tempfile
from typing import Literal, Optional, List
from pathlib import Path
from dotenv import load_dotenv

# Load .env từ thư mục gốc của app (ct_datalake/)
_env_path = Path(__file__).resolve().parent / ".env"
if not _env_path.exists():
    _env_path = Path(__file__).resolve().parents[1] / ".env"
load_dotenv(dotenv_path=_env_path, override=True)

# ─── internal modules ───────────────────────────────────────────────
from .search import search as faiss_search
from .llm_rerank import rerank
from .jd_match import (
    match_jd,
    normalize_candidate,
    parse_jd,
    extract_text_from_upload as _extract_text,
)
from .ai_matching import extract_keywords, search_domain
from openai import OpenAI
from .openai_key import get_openai_client
from . import ai_matching as _rc

# ─── OpenAI client ──────────────────────────────────────────────────

# ════════════════════════════════════════════════════════════════════
# ROUTES (Frappe Whitelisted)
# ════════════════════════════════════════════════════════════════════

# ── Health check ────────────────────────────────────────────────────
@frappe.whitelist(allow_guest=False)
def root():
    """Kiểm tra trạng thái API"""
    return {
        "status": "ok",
        "service": "AI Candidate Search API (Frappe)",
        "version": "1.0.0",
        "endpoints": [
            "/api/method/ct_datalake.api.root",
            "/api/method/ct_datalake.api.health",
            "/api/method/ct_datalake.api.semantic_search",
            "/api/method/ct_datalake.api.semantic_search_llm",
            "/api/method/ct_datalake.api.jd_match",
            "/api/method/ct_datalake.api.jd_match_upload",
            "/api/method/ct_datalake.api.g600_analyze",
        ],
    }

@frappe.whitelist(allow_guest=False)
def health():
    """Health check chi tiết"""
    return {
        "status": "healthy",
        "openai_configured": bool(get_openai_client()),
    }

# ── Semantic Search (FAISS) ─────────────────────────────────────────
@frappe.whitelist(allow_guest=False)
def semantic_search(query: str, mode: str = "in", top_k: int = 5):
    """
    Tìm kiếm ứng viên sử dụng FAISS semantic search.
    """
    try:
        top_k = int(top_k)
        raw = faiss_search(query=query, mode=mode, top_k=top_k)
    except Exception as e:
        frappe.throw(f"Search error: {str(e)}")

    results = []
    for r in raw:
        norm = normalize_candidate(r["data"], mode)
        results.append({
            "faiss_score": round(r["score"], 4),
            "candidate": norm,
            "raw": r["data"],
        })

    return {
        "query": query,
        "mode": mode,
        "total": len(results),
        "results": results,
    }

# ── Semantic Search (LLM Rerank) ────────────────────────────────────
@frappe.whitelist(allow_guest=False)
def semantic_search_llm(query: str, mode: str = "in", top_k: int = 5):
    """
    Tìm kiếm FAISS rồi đưa kết quả vào LLM để phân tích.
    """
    if not get_openai_client():
        frappe.throw("OPENAI_API_KEY chưa được cấu hình")

    try:
        top_k = int(top_k)
        raw = faiss_search(
            query=query,
            mode=mode,
            top_k=max(top_k * 2, 10)
        )
    except Exception as e:
        frappe.throw(f"Search error: {str(e)}")

    if not raw:
        return {"query": query, "mode": mode, "total_found": 0, "llm_analysis": ""}

    try:
        analysis = rerank(query=query, candidates=raw[:top_k], mode=mode)
    except Exception as e:
        frappe.throw(f"LLM rerank error: {str(e)}")

    return {
        "query": query,
        "mode": mode,
        "total_found": len(raw),
        "candidates_sent_to_llm": top_k,
        "llm_analysis": analysis,
    }

# ── JD Matching (via form fields) ───────────────────────────────────
@frappe.whitelist(allow_guest=False)
def jd_match(jd_text: str, mode: str = "in", top_k: int = 5, fast: bool = False):
    """
    Match JD với ứng viên.
    """
    if not get_openai_client() and not frappe.parse_json(fast):
        frappe.throw("OPENAI_API_KEY chưa cấu hình. Dùng fast=true để chỉ dùng FAISS.")

    try:
        top_k = int(top_k)
        fast = frappe.parse_json(fast)
        result = match_jd(
            jd_text=jd_text,
            mode=mode,
            top_k=top_k,
            fast=fast,
        )
    except Exception as e:
        frappe.throw(f"JD match error: {str(e)}")

    return {
        "parsed_jd": result.get("parsed_jd", {}),
        "total_found": result.get("total_found", 0),
        "candidates": result.get("candidates", []),
    }

# ── JD Matching (File Upload) ───────────────────────────────────────
@frappe.whitelist(allow_guest=False)
def jd_match_upload():
    """
    Upload file JD → extract text → match ứng viên.
    Sử dụng frappe.request.files
    """
    if not frappe.request.files:
        frappe.throw("Chưa có file nào được upload")

    file = frappe.request.files.get("file")
    if not file:
        frappe.throw("Vui lòng upload file với key 'file'")

    mode = frappe.form_dict.get("mode", "in")
    top_k = int(frappe.form_dict.get("top_k", 5))
    fast = frappe.parse_json(frappe.form_dict.get("fast", "false"))

    allowed = (".txt", ".md", ".pdf", ".docx")
    name = (file.filename or "").lower()
    if not any(name.endswith(ext) for ext in allowed):
        frappe.throw(f"Định dạng không hỗ trợ. Chỉ chấp nhận: {', '.join(allowed)}")

    # Tạo mock object để tái dùng hàm extract_text_from_upload
    class _MockUpload:
        def __init__(self, name, data):
            self.name = name
            self._data = data
        def read(self):
            return self._data

    try:
        jd_text = _extract_text(_MockUpload(file.filename, file.stream.read()))
    except Exception as e:
        frappe.throw(str(e))

    try:
        result = match_jd(
            jd_text=jd_text,
            mode=mode,
            top_k=top_k,
            fast=fast,
        )
    except Exception as e:
        frappe.throw(f"JD match error: {str(e)}")

    return {
        "parsed_jd": result.get("parsed_jd", {}),
        "total_found": result.get("total_found", 0),
        "candidates": result.get("candidates", []),
    }

# ── G600 PDF Analysis ────────────────────────────────────────────────
@frappe.whitelist(allow_guest=False)
def g600_analyze():
    """
    Upload PDF tờ trình G600 → GPT-4o Vision đọc và trích xuất lĩnh vực.
    """
    if not get_openai_client():
        frappe.throw("OPENAI_API_KEY chưa được cấu hình")

    if not frappe.request.files:
        frappe.throw("Chưa có file nào được upload")

    file = frappe.request.files.get("file")
    if not file:
        frappe.throw("Vui lòng upload file với key 'file'")

    source = frappe.form_dict.get("source", "both")
    top_k = int(frappe.form_dict.get("top_k", 5))
    score_threshold = float(frappe.form_dict.get("score_threshold", 0.40))

    if not (file.filename or "").lower().endswith(".pdf"):
        frappe.throw("Chỉ chấp nhận file PDF")

    file_bytes = file.stream.read()

    # Ghi PDF tạm thời
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name

    try:
        domains = extract_keywords(tmp_path, get_openai_client())
    except Exception as e:
        frappe.throw(f"GPT Vision error: {str(e)}")
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)

    if not domains:
        frappe.throw("GPT không trích xuất được lĩnh vực nào từ PDF")

    datasets = get_datasets()
    embed_model = get_embed_model()

    if source == "in":
        active = {k: v for k, v in datasets.items() if k == "in"}
    elif source == "out":
        active = {k: v for k, v in datasets.items() if k == "out"}
    else:
        active = datasets

    if not active:
        frappe.throw("Không tìm thấy FAISS index. Kiểm tra file index.faiss / index_out.faiss")

    _rc.TOP_K = top_k
    _rc.SCORE_THRESHOLD = score_threshold

    domain_results = []
    for domain in domains:
        hits = search_domain(domain, active, embed_model)
        candidates = []
        for hit in hits:
            d = hit["data"]
            mode = hit["mode"]
            norm = normalize_candidate(d, mode)
            candidates.append({
                "source": "🌐 Google Scholar" if mode == "out" else "🏫 Nội bộ",
                "faiss_score": hit["score"],
                "candidate": norm,
                "raw": d,
            })

        domain_results.append({
            "domain_name": domain["name"],
            "query_vi": domain.get("query_vi", ""),
            "query_en": domain.get("query_en", ""),
            "keywords_vi": domain.get("keywords_vi", []),
            "keywords_en": domain.get("keywords_en", []),
            "total_candidates": len(candidates),
            "candidates": candidates,
        })

    return {
        "source": source,
        "top_k_per_domain": top_k,
        "score_threshold": score_threshold,
        "total_domains": len(domain_results),
        "domains": domain_results,
    }

# ── Parse JD only ────────────────────────────────────────────────────
@frappe.whitelist(allow_guest=False)
def jd_parse_only(query: str):
    """
    Gửi JD text → GPT phân tích.
    """
    if not get_openai_client():
        frappe.throw("OPENAI_API_KEY chưa cấu hình")
    try:
        parsed = parse_jd(query)
    except Exception as e:
        frappe.throw(str(e))
    return {"parsed_jd": parsed}
# ── Draft Document ────────────────────────────────────────────────────────────
@frappe.whitelist(allow_guest=False)
def draft_document(
    candidate_info: str,
    doc_type: str = "invite_collab",
    org_name: str = "",
    sender_name: str = "",
    extra_note: str = "",
):
    """
    Soạn thảo văn bản mời ứng viên dựa trên thông tin ứng viên.
    
    Params:
        candidate_info: JSON string chứa thông tin ứng viên (name, school, expertise, ...)
        doc_type: Loại văn bản (invite_collab, invite_expert, invite_project, invite_lecture, consult_request, partnership)
        org_name: Tên đơn vị gửi
        sender_name: Người ký
        extra_note: Ghi chú thêm
    """
    if not get_openai_client():
        frappe.throw("OPENAI_API_KEY chưa được cấu hình")

    DRAFT_PROMPTS = {
        "invite_collab":   "Thư mời hợp tác nghiên cứu khoa học",
        "invite_expert":   "Thư mời tham gia hội đồng chuyên gia phản biện",
        "invite_project":  "Thư mời tham gia dự án nghiên cứu",
        "invite_lecture":  "Thư mời giảng dạy hoặc báo cáo chuyên đề",
        "consult_request": "Công văn đề nghị tư vấn chuyên môn",
        "partnership":     "Thư đề xuất hợp tác chiến lược dài hạn",
    }

    doc_label = DRAFT_PROMPTS.get(doc_type, "Thư mời hợp tác")
    org = org_name or "đơn vị chúng tôi"
    sender_line = f"Người ký: {sender_name}" if sender_name else ""
    extra_line = f"Lưu ý thêm: {extra_note}" if extra_note else ""

    prompt = f"""Bạn là chuyên viên soạn thảo văn bản hành chính - ngoại giao chuyên nghiệp.
Hãy soạn một "{doc_label}" bằng tiếng Việt, trang trọng, chuyên nghiệp và đầy đủ.

Thông tin ứng viên/chuyên gia cần mời:
{candidate_info}

Đơn vị gửi: {org}
{sender_line}
{extra_line}

Yêu cầu:
- Văn phong lịch sự, trang trọng, đúng phong cách văn bản hành chính Việt Nam
- Đề cập cụ thể đến lĩnh vực chuyên môn của ứng viên
- Có đầy đủ: Kính gửi, Nội dung chính, Lời kết, Ký tên
- Độ dài phù hợp (khoảng 200-350 từ)
- Để trống [ngày tháng], [địa điểm], [số điện thoại liên hệ] nếu chưa có thông tin"""

    try:
        response = get_openai_client().chat.completions.create(
            model="gpt-4o",
            messages=[{"role": "user", "content": prompt}],
            max_tokens=1000,
            temperature=0.7,
        )
        draft_text = response.choices[0].message.content
    except Exception as e:
        frappe.throw(f"GPT error: {str(e)}")

    return {
        "doc_type": doc_type,
        "doc_label": doc_label,
        "draft": draft_text,
    }


# ── AI Candidate Persona Report ──────────────────────────────────
@frappe.whitelist(allow_guest=True)
def generate_candidate_report(
    ai_test_url: str,
    g5_test_url: str,
    eq_test_url: str,
    survey_url: str = "",
    jd_text: str = "",
):
    """
    Pipeline tạo AI Candidate Persona Report.

    Params (tất cả đều là string, không hard-code):
        ai_test_url  : Link Odoo survey print – bài test AI
        g5_test_url  : Link Odoo survey print – bài test 5G
        eq_test_url  : Link jobtest.vn – bài EQ/IQ
        survey_url   : Link phiếu đánh giá phỏng vấn (có thể trống)
        jd_text      : Nội dung JD (string)

    CV được upload qua frappe.request.files['cv_file'] (pdf/docx).

    Returns: dict chứa toàn bộ fields + doctype_name đã lưu.
    """
    import warnings, requests
    import PyPDF2, io
    from bs4 import BeautifulSoup
    from docx import Document as DocxDocument
    from pathlib import Path

    warnings.filterwarnings("ignore")

    _client = get_openai_client()
    if not _client:
        frappe.throw("OPENAI_API_KEY chưa được cấu hình")

    BASE_DIR = Path(__file__).resolve().parent.parent / "AI_ATS"
    # fallback nếu không tìm thấy
    if not BASE_DIR.exists():
        BASE_DIR = Path("/Users/_qh.fol_/AI_ATS")

    AI_SCORING_PDF  = BASE_DIR / "CTG-KNC-TD-QĐ04.BM02-HƯỚNG DẪN CHẤM ĐIỂM BÀI TEST NĂNG LỰC AI (1).pdf"
    SWAT_PRD_DOCX   = BASE_DIR / "18052026_RD - PRD - AI Candidate Persona Report.docx"
    G5_SCORING_DOCX = BASE_DIR / "CTG-KNC-TD-QT01.BM16 - Bộ CÂU HỎI ĐÁNH GIÁ TIỀM NĂNG ỨNG VIÊN 2 1 (2).docx"

    # ── Helpers ──────────────────────────────────────────────────────
    def _read_pdf_bytes(data: bytes) -> str:
        reader = PyPDF2.PdfReader(io.BytesIO(data))
        return "\n".join(pg.extract_text() or "" for pg in reader.pages).strip()

    def _read_docx_bytes(data: bytes) -> str:
        doc = DocxDocument(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for i, tbl in enumerate(doc.tables):
            parts.append(f"[TABLE {i+1}]")
            for row in tbl.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts).strip()

    def _read_pdf_path(path: Path) -> str:
        return _read_pdf_bytes(path.read_bytes()) if path.exists() else ""

    def _read_docx_path(path: Path) -> str:
        return _read_docx_bytes(path.read_bytes()) if path.exists() else ""

    def _scrape(url: str) -> str:
        if not url:
            return "[Chưa có]"
        try:
            resp = requests.get(url, timeout=20, verify=False)
            soup = BeautifulSoup(resp.text, "html.parser")
            for tag in soup(["script", "style", "nav", "footer", "header"]):
                tag.decompose()
            text = soup.get_text(separator="\n", strip=True)
            return text if len(text) > 100 else f"[Nội dung rỗng: {url}]"
        except Exception as e:
            return f"[Lỗi scrape: {e}]"

    # ── Đọc CV từ upload ─────────────────────────────────────────────
    cv_text = ""
    if frappe.request.files:
        cv_file = frappe.request.files.get("cv_file")
        if cv_file:
            raw = cv_file.stream.read()
            fname = (cv_file.filename or "").lower()
            if fname.endswith(".pdf"):
                cv_text = _read_pdf_bytes(raw)
            elif fname.endswith(".docx"):
                cv_text = _read_docx_bytes(raw)

    # ── Scrape test links ────────────────────────────────────────────
    ai_test_text = _scrape(ai_test_url)
    g5_test_text = _scrape(g5_test_url)
    eq_test_text = _scrape(eq_test_url)
    survey_text  = _scrape(survey_url) if survey_url else "[Chưa có]"

    # ── Internal scoring guides ──────────────────────────────────────
    ai_scoring = _read_pdf_path(AI_SCORING_PDF)
    swat_prd   = _read_docx_path(SWAT_PRD_DOCX)
    g5_scoring = _read_docx_path(G5_SCORING_DOCX)

    # ── Build prompts ────────────────────────────────────────────────
    SYSTEM = """Bạn là AI Agent đánh giá ứng viên cho CT Group (NoAI-NoHire).
Áp dụng ĐÚNG 2 bộ tiêu chí:

Bộ 1 – ĐIỂM BÀI TEST AI (10 câu × 10đ = 100đ)
Chấm từng câu theo AI_SCORING_GUIDE.
Ngưỡng: ≥75 Xuất sắc | 60-74 Khá (AI-Ready) | 50-59 Theo dõi | <50 Non-AI

Bộ 2 – SWAT ELITE (4 trụ cột, thang 0-10, trọng số 50/20/20/10%)
Tổng trọng số ≥6.0/10 → SWAT Elite ĐẠT

CHỈ trả về JSON hợp lệ theo schema sau:
{
  "candidate_name": "string",
  "position": "string",
  "email": "string",
  "phone": "string",
  "ai_test_table": [
    {"cau":1,"noi_dung":"AI Awareness","diem_toi_da":10,"diem_cham":0,"ly_do":"string"},
    {"cau":2,"noi_dung":"AI Daily Use","diem_toi_da":10,"diem_cham":0,"ly_do":"string"},
    {"cau":3,"noi_dung":"AI Self-Assessment","diem_toi_da":10,"diem_cham":0,"ly_do":"string"},
    {"cau":4,"noi_dung":"AI Problem Solving","diem_toi_da":10,"diem_cham":0,"ly_do":"string"},
    {"cau":5,"noi_dung":"Prompt Engineering","diem_toi_da":10,"diem_cham":0,"ly_do":"string"},
    {"cau":6,"noi_dung":"AI x Teamwork","diem_toi_da":10,"diem_cham":0,"ly_do":"string"},
    {"cau":7,"noi_dung":"AI Productivity","diem_toi_da":10,"diem_cham":0,"ly_do":"string"},
    {"cau":8,"noi_dung":"AI Mindset","diem_toi_da":10,"diem_cham":0,"ly_do":"string"},
    {"cau":9,"noi_dung":"AI Limitation","diem_toi_da":10,"diem_cham":0,"ly_do":"string"},
    {"cau":10,"noi_dung":"AI Growth Plan","diem_toi_da":10,"diem_cham":0,"ly_do":"string"}
  ],
  "ai_test_total": 0,
  "ai_test_label": "AI-Ready",
  "swat_table": [
    {"tru_cot":"AI First Mindset","ty_trong":"50%","diem_tho":0,"diem_trong_so":0.0,"co_so":"string"},
    {"tru_cot":"2AS Execution Capability","ty_trong":"20%","diem_tho":0,"diem_trong_so":0.0,"co_so":"string"},
    {"tru_cot":"Practical Efficiency & Productivity","ty_trong":"20%","diem_tho":0,"diem_trong_so":0.0,"co_so":"string"},
    {"tru_cot":"Risk Control & Language","ty_trong":"10%","diem_tho":0,"diem_trong_so":0.0,"co_so":"string"}
  ],
  "swat_total": 0.0,
  "swat_label": "ĐẠT",
  "strengths": "string",
  "gaps": "string",
  "best_at": "string",
  "conclusion": "string",
  "next_steps": "string",
  "decision": "ĐẠT"
}"""

    from datetime import datetime
    USER = f"""
### CV:
{cv_text[:3000]}

### JD:
{jd_text[:1500]}

### HƯỚNG DẪN CHẤM AI TEST:
{ai_scoring[:2500]}

### FRAMEWORK SWAT ELITE:
{swat_prd[:1500]}

### TIÊU CHÍ 5G:
{g5_scoring[:800]}

### KẾT QUẢ TEST AI (link):
{ai_test_text[:2500]}

### KẾT QUẢ TEST 5G (link):
{g5_test_text[:2500]}

### KẾT QUẢ EQ/IQ (link):
{eq_test_text[:1500]}

### PHIẾU PHỎNG VẤN:
{survey_text[:1000]}

Ngày: {datetime.now().strftime("%d/%m/%Y %H:%M")}
Trả về JSON hợp lệ, điền đủ mọi trường."""

    # ── Gọi OpenAI ───────────────────────────────────────────────────
    resp = _client.chat.completions.create(
        model="gpt-4o",
        messages=[
            {"role": "system", "content": SYSTEM},
            {"role": "user",   "content": USER},
        ],
        response_format={"type": "json_object"},
        temperature=0.2,
        max_tokens=4000,
    )
    data = json.loads(resp.choices[0].message.content)

    # ── Lưu vào DocType ──────────────────────────────────────────────
    doc = frappe.get_doc({
        "doctype": "AI Candidate Report",
        # --- text fields ---
        "candidate_name": data.get("candidate_name", "Unknown"),
        "position":       data.get("position", ""),
        "email":          data.get("email", ""),
        "phone":          data.get("phone", ""),
        "analysis_date":  datetime.now(),
        "ai_test_url":    ai_test_url,
        "g5_test_url":    g5_test_url,
        "eq_test_url":    eq_test_url,
        "survey_url":     survey_url,
        "jd_text":        jd_text[:2000],
        "ai_test_total":  data.get("ai_test_total", 0),
        "ai_test_label":  data.get("ai_test_label", ""),
        "swat_total":     data.get("swat_total", 0),
        "swat_label":     data.get("swat_label", ""),
        "strengths":      data.get("strengths", ""),
        "gaps":           data.get("gaps", ""),
        "best_at":        data.get("best_at", ""),
        "conclusion":     data.get("conclusion", ""),
        "next_steps":     data.get("next_steps", ""),
        "decision":       data.get("decision", ""),
        # --- JSON text fields (2 bảng) ---
        "ai_test_table":  json.dumps(data.get("ai_test_table", []), ensure_ascii=False),
        "swat_table":     json.dumps(data.get("swat_table",    []), ensure_ascii=False),
    })
    doc.insert(ignore_permissions=True)
    frappe.db.commit()

    return {
        "doctype_name": doc.name,
        **data,
    }


# ════════════════════════════════════════════════════════════════════
# CT Frappe Template — Session & Access Control
# ════════════════════════════════════════════════════════════════════

from ct_datalake.utils.activity_logger import ActivityLogger, Timer

_logger = ActivityLogger(prefix="DL", module="CT DataLake")


def _resolve_session(session_id: str) -> str:
    """Tìm session_name từ session_id. Fallback trả về chuỗi rỗng."""
    if not session_id:
        return ""
    try:
        rows = frappe.db.get_all(
            "DL Session",
            filters={"session_id": session_id},
            fields=["name"],
            limit=1,
            ignore_permissions=True,  # bắt buộc: bypass DocType read permission
        )
        return rows[0].name if rows else ""
    except Exception:
        return ""


@frappe.whitelist(allow_guest=False)
def get_context():
    """
    Entry point cho Frontend (initSession).
    - Xác thực quyền qua ct_agent_hub.check_app_access (cookie-based)
    - Tạo DL Session mới
    - Trả về csrf_token + session_id + user info
    """
    dept = ""
    role = ""
    try:
        from ct_agent_hub.api import check_app_access
        agents_data = check_app_access("ct_datalake")
        user_depts = agents_data.get("user_departments", [])
        dept = ",".join(user_depts) if user_depts else ""
        role = agents_data.get("user_role", "")
    except ImportError:
        pass  # ct_agent_hub chưa được cài đặt

    session_id   = str(uuid.uuid4())
    session_name = _logger.create_session(session_id, dept=dept, role=role)

    return {
        "csrf_token":   frappe.sessions.get_csrf_token(),
        "session_id":   session_id,
        "session_name": session_name,
        "user":         frappe.session.user,
        "full_name":    frappe.utils.get_fullname(frappe.session.user),
    }

